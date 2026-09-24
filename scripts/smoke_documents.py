"""Synthetic documents shared by browser and API smoke tests; no private corpus needed."""

import base64
import json
from io import BytesIO
from uuid import uuid4

from docx import Document
from reportlab.pdfgen.canvas import Canvas


def make_documents():
    identity = str(uuid4())
    pdf_bytes = BytesIO()
    pdf = Canvas(pdf_bytes, invariant=1)
    pdf.setSubject(identity)
    pdf.drawString(60, 740, "Alpha retention policy")
    pdf.drawString(60, 715, "Alpha retains records for 30 days.")
    pdf.drawString(60, 690, "Alpha executes requests synchronously.")
    pdf.save()
    docx_bytes = BytesIO()
    document = Document()
    document.core_properties.identifier = identity
    document.add_heading("Beta retention policy", level=1)
    document.add_paragraph("Beta retains records for 90 days.")
    document.add_paragraph("Beta executes requests asynchronously.")
    document.save(docx_bytes)
    return [
        ("alpha-verification.pdf", pdf_bytes.getvalue()),
        ("beta-verification.docx", docx_bytes.getvalue()),
    ]


if __name__ == "__main__":
    print(
        json.dumps(
            [
                {"name": name, "base64": base64.b64encode(content).decode()}
                for name, content in make_documents()
            ]
        )
    )
