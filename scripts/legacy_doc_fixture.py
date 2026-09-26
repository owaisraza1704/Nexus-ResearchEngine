"""Generate a synthetic binary DOC fixture inside the API image; print a JSON/base64 payload.

docker compose exec -T api python < scripts/legacy_doc_fixture.py
LibreOffice, not custom code, writes the legacy Word format.
"""

import base64
import json
import subprocess
from pathlib import Path
from tempfile import TemporaryDirectory
from uuid import uuid4

from docx import Document

with TemporaryDirectory() as directory:
    root = Path(directory)
    document = Document()
    document.core_properties.identifier = uuid4().hex
    document.add_heading("Gamma legacy retention policy", level=1)
    document.add_paragraph("Gamma retains records for 60 days.")
    document.add_paragraph("Gamma executes requests asynchronously.")
    source = root / "gamma-verification.docx"
    document.save(source)
    subprocess.run(
        [
            "soffice",
            f"-env:UserInstallation={(root / 'profile').as_uri()}",
            "--headless",
            "--convert-to",
            "doc:MS Word 97",
            "--outdir",
            str(root),
            str(source),
        ],
        check=True,
        capture_output=True,
        timeout=45,
    )
    content = source.with_suffix(".doc").read_bytes()
    assert content.startswith(bytes.fromhex("d0cf11e0a1b11ae1"))
    print(
        json.dumps({"name": "gamma-verification.doc", "base64": base64.b64encode(content).decode()})
    )
