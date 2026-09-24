from io import BytesIO
from unittest.mock import Mock
from uuid import UUID, uuid4

import pytest
from docx import Document as WordDocument
from fastapi.testclient import TestClient
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.db.models import ChunkEmbedding, Document, DocumentChunk, Source
from app.db.session import get_db
from app.ingestion.docling_parser import DocumentHasNoText
from app.main import create_app


@pytest.fixture
def word_bytes():
    document = WordDocument()
    document.add_heading("Research", level=1)
    document.add_paragraph("Nexus produces cited research answers.")
    document.add_heading("Evidence", level=1)
    document.add_paragraph("Sources retain document versions.")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_upload_docx_persists_and_exposes_document_chunks(
    client,
    db,
    azure_api,
    word_bytes,
):
    response = client.post(
        "/v1/sources/uploads",
        files={
            "file": ("research.docx", word_bytes),
        },
        data={"display_name": "Research notes"},
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["status"] == "ready"
    assert body["document"]["chunk_count"] >= 2
    source_id = body["source_id"]
    detail = client.get(f"/v1/sources/{source_id}").json()
    assert "Nexus produces cited research answers." in detail["normalized_text"]
    assert detail["document"]["page_count"] is None
    chunks = client.get(f"/v1/sources/{source_id}/chunks").json()
    assert chunks["total"] == body["document"]["chunk_count"]
    assert [chunk["sequence"] for chunk in chunks["chunks"]] == list(range(chunks["total"]))
    assert all(chunk["char_count"] == len(chunk["text"]) for chunk in chunks["chunks"])
    assert all(chunk["locator"]["doc_items"] for chunk in chunks["chunks"])
    assert db.scalar(select(func.count()).select_from(ChunkEmbedding)) == chunks["total"]
    document = db.get(Document, UUID(body["document"]["document_id"]))
    assert document.ready_at is not None
    assert document.document_metadata["chunker_name"] == "docling.HierarchicalChunker"
    assert document.document_metadata["ingestion_duration_ms"] > 0
    listing = client.get("/v1/sources").json()
    assert listing["total"] == 1
    assert listing["sources"][0]["display_name"] == "Research notes"
    assert listing["sources"][0]["chunk_count"] == chunks["total"]
    assert len(azure_api["calls"]) == 1


def test_failed_upload_can_retry_and_ready_duplicate_is_rejected(
    client,
    db,
    monkeypatch,
    azure_api,
    word_bytes,
):
    from app.api import sources

    original_parser = sources.parse_document

    def fail_parse(*args, **kwargs):
        raise DocumentHasNoText("The document contains no extractable text")

    monkeypatch.setattr(sources, "parse_document", fail_parse)
    failure = client.post("/v1/sources/uploads", files={"file": ("research.docx", word_bytes)})
    assert failure.status_code == 422
    source = db.scalar(select(Source))
    source_id = source.id
    assert source.status == "failed"
    assert source.error_code == "DOCUMENT_HAS_NO_TEXT"
    assert db.scalar(select(func.count()).select_from(Document)) == 0

    monkeypatch.setattr(sources, "parse_document", original_parser)
    retry = client.post("/v1/sources/uploads", files={"file": ("research.docx", word_bytes)})
    assert retry.status_code == 201, retry.text
    assert retry.json()["source_id"] == str(source_id)
    db.refresh(source)
    assert source.error_code is None
    duplicate = client.post("/v1/sources/uploads", files={"file": ("research.docx", word_bytes)})
    assert duplicate.status_code == 409
    assert duplicate.json()["error"]["code"] == "SOURCE_ALREADY_EXISTS"
    assert db.scalar(select(func.count()).select_from(Source)) == 1
    assert db.scalar(select(func.count()).select_from(Document)) == 1
    assert len(azure_api["calls"]) == 1


@pytest.mark.parametrize(
    ("name", "content", "code", "http_status"),
    [
        ("empty.pdf", b"", "EMPTY_FILE", 422),
        ("text.txt", b"text", "UNSUPPORTED_MEDIA_TYPE", 415),
        ("large.pdf", b"01234567890", "FILE_TOO_LARGE", 413),
    ],
)
def test_upload_rejects_input_before_database(settings, name, content, code, http_status):
    settings.max_upload_bytes = 10
    database = Mock(spec=Session)
    app = create_app()
    app.dependency_overrides[get_settings] = lambda: settings
    app.dependency_overrides[get_db] = lambda: database
    response = TestClient(app).post("/v1/sources/uploads", files={"file": (name, content)})
    assert response.status_code == http_status
    assert response.json()["error"]["code"] == code
    assert database.mock_calls == []


@pytest.mark.parametrize(
    ("setting", "value", "code"),
    [
        ("max_document_chars", 10, "DOCUMENT_TOO_LARGE"),
        ("max_document_chunks", 1, "TOO_MANY_CHUNKS"),
        ("max_chunk_chars", 10, "CHUNK_TOO_LARGE"),
    ],
)
def test_ingestion_limits_fail_before_embeddings(
    client,
    db,
    settings,
    azure_api,
    word_bytes,
    setting,
    value,
    code,
):
    setattr(settings, setting, value)
    response = client.post("/v1/sources/uploads", files={"file": ("research.docx", word_bytes)})
    assert response.status_code == 422, response.text
    assert response.json()["error"]["code"] == code
    assert azure_api["calls"] == []
    assert db.scalar(select(Source)).status == "failed"
    assert db.scalar(select(func.count()).select_from(DocumentChunk)) == 0


def test_inspection_needs_no_provider_and_enforces_document_ownership(
    client,
    ready_source,
    settings,
    azure_api,
):
    source, _, _ = ready_source()
    other, other_document, _ = ready_source()
    settings.azure_openai_api_key = None
    assert client.get("/v1/sources").status_code == 200
    assert client.get(f"/v1/sources/{source.id}").status_code == 200
    assert (
        client.get(f"/v1/sources/{source.id}/chunks?limit=1&offset=1").json()["chunks"][0][
            "sequence"
        ]
        == 1
    )
    assert (
        client.get(f"/v1/sources/{source.id}/chunks?document_id={other_document.id}").status_code
        == 404
    )
    assert client.get(f"/v1/sources/{uuid4()}").status_code == 404
    assert client.get("/v1/sources?limit=101").status_code == 422
    assert azure_api["calls"] == []


def test_ingestion_budget_prevents_embedding_call(
    client,
    db,
    word_bytes,
    azure_api,
    monkeypatch,
):
    from app.api import sources

    monkeypatch.setattr(sources, "perf_counter", Mock(side_effect=[0.0, 181.0]))
    response = client.post("/v1/sources/uploads", files={"file": ("research.docx", word_bytes)})
    assert response.status_code == 504
    assert response.json()["error"]["code"] == "INGESTION_TIMEOUT"
    assert db.scalar(select(Source)).status == "failed"
    assert azure_api["calls"] == []
