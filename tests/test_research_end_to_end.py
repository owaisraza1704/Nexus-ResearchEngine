from io import BytesIO

from docx import Document
from reportlab.pdfgen.canvas import Canvas


def test_pdf_docx_upload_to_research_and_persisted_reload(client, research_azure):
    pdf_bytes = BytesIO()
    pdf = Canvas(pdf_bytes, invariant=1)
    pdf.drawString(60, 740, "Alpha retains records for 30 days.")
    pdf.save()
    docx_bytes = BytesIO()
    document = Document()
    document.add_heading("Beta retention policy", level=1)
    document.add_paragraph("Beta retains records for 90 days.")
    document.save(docx_bytes)

    sources = []
    for name, content in (
        ("alpha.pdf", pdf_bytes.getvalue()),
        ("beta.docx", docx_bytes.getvalue()),
    ):
        response = client.post("/v1/sources/uploads", files={"file": (name, content)})
        assert response.status_code == 201, response.text
        assert response.json()["status"] == "ready"
        sources.append(response.json()["source_id"])

    response = client.post(
        "/v1/research/runs",
        json={
            "question": "How do the retention policies differ?",
            "source_ids": sources,
            "retrieval": {"top_k_per_source": 1},
        },
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "completed"
    assert len(body["source_coverage"]) == 2
    for evidence in body["evidence"]:
        chunks_response = client.get(
            f"/v1/sources/{evidence['source_id']}/chunks",
            params={
                "document_id": evidence["document_id"],
            },
        )
        chunks = {chunk["chunk_id"]: chunk for chunk in chunks_response.json()["chunks"]}
        assert evidence["locator"] == chunks[evidence["chunk_id"]]["locator"]
        assert evidence["excerpt"] == chunks[evidence["chunk_id"]]["text"]
    assert "page(s) 1" in body["evidence"][0]["display_text"]
    assert "Beta retention policy" in body["evidence"][1]["display_text"]
    assert client.get(f"/v1/research/runs/{body['run_id']}/result").json() == body
