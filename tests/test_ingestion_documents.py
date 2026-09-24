from io import BytesIO

import pytest
from docx import Document
from PIL import Image, ImageDraw
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen.canvas import Canvas

from app.ingestion.docling_chunker import chunk_document
from app.ingestion.docling_parser import DocumentHasNoText, DocumentParseError, parse_document


def test_docx_normalization_chunks_and_locators_are_repeatable(tmp_path):
    path = tmp_path / "research.docx"
    document = Document()
    document.add_heading("Project objective", level=1)
    document.add_paragraph("Nexus produces source-grounded answers with citations.")
    document.add_heading("Evidence", level=1)
    document.add_paragraph("Source locators let the reader inspect the evidence.")
    document.save(path)
    first = parse_document(path)
    second = parse_document(path)
    assert first.normalized_text_sha256 == second.normalized_text_sha256
    assert chunk_document(first) == chunk_document(second)
    assert any(chunk.locator.get("headings") for chunk in chunk_document(first))


def test_pdf_upload_locators_and_repeatability(client, azure_api, tmp_path):
    path = tmp_path / "research.pdf"
    pdf = Canvas(str(path), invariant=1)
    pdf.drawString(60, 740, "Nexus produces source-grounded answers with citations.")
    pdf.showPage()
    pdf.drawString(60, 740, "Each citation points to evidence in a source document.")
    pdf.save()
    first = parse_document(path)
    second = parse_document(path)
    assert first.page_count == 2
    assert first.normalized_text_sha256 == second.normalized_text_sha256
    assert chunk_document(first) == chunk_document(second)
    assert {block.locator["page"] for block in first.blocks} == {1, 2}
    with path.open("rb") as upload:
        response = client.post("/v1/sources/uploads", files={"file": (path.name, upload)})
    assert response.status_code == 201, response.text
    assert response.json()["document"]["page_count"] == 2
    source_id = response.json()["source_id"]
    chunks = client.get(f"/v1/sources/{source_id}/chunks").json()["chunks"]
    pages = {
        prov["page_no"]
        for chunk in chunks
        for item in chunk["locator"]["doc_items"]
        for prov in item.get("prov", [])
    }
    assert pages == {1, 2}
    with pytest.raises(DocumentParseError):
        parse_document(path, max_pages=1)


def test_image_only_pdf_does_not_silently_use_ocr(tmp_path):
    path = tmp_path / "scanned.pdf"
    image = Image.new("RGB", (400, 100), "white")
    ImageDraw.Draw(image).text((10, 20), "Scanned research note", fill="black")
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    pdf = Canvas(str(path))
    pdf.drawImage(ImageReader(buffer), 60, 600, width=400, height=100)
    pdf.save()
    with pytest.raises(DocumentHasNoText):
        parse_document(path)


@pytest.mark.parametrize("suffix", [".pdf", ".docx"])
def test_malformed_document_fails_explicitly(tmp_path, suffix):
    path = tmp_path / f"malformed{suffix}"
    path.write_bytes(b"This is not a document file")
    with pytest.raises(DocumentParseError):
        parse_document(path)
