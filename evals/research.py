"""Exercise the real research API with synthetic documents and inspectable evidence."""

import argparse
import json
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from time import perf_counter
from xml.sax.saxutils import escape

import httpx
from docx import Document
from ir_measures import RR, Qrel, R, ScoredDoc, calc, nDCG
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from app.api.research import ResearchResponse

DEFAULT_CASES = Path(__file__).with_name("research_cases.json")


def document_bytes(spec: dict) -> bytes:
    """Let python-docx and ReportLab create the two formats our ingestion already supports."""
    buffer = BytesIO()
    if spec["filename"].endswith(".docx"):
        document = Document()
        document.add_heading(spec["title"], level=0)
        for section in spec["sections"]:
            document.add_heading(section["heading"], level=1)
            document.add_paragraph(section["text"])
        document.save(buffer)
    else:
        styles = getSampleStyleSheet()
        story = [Paragraph(escape(spec["title"]), styles["Title"])]
        for section in spec["sections"]:
            story.extend(
                [
                    Paragraph(escape(section["heading"]), styles["Heading1"]),
                    Paragraph(escape(section["text"]), styles["BodyText"]),
                    Spacer(1, 12),
                ]
            )
        SimpleDocTemplate(buffer, invariant=1).build(story)
    return buffer.getvalue()


def upload_documents(client: httpx.Client, dataset: dict) -> dict:
    sources = {}
    for key, spec in dataset["documents"].items():
        started = perf_counter()
        response = client.post(
            "/v1/sources/uploads", files={"file": (spec["filename"], document_bytes(spec))}
        )
        response.raise_for_status()
        body = response.json()
        sources[key] = {
            "source_id": body["source_id"],
            "document_id": body["document"]["document_id"],
            "filename": spec["filename"],
            "ingestion_ms": round((perf_counter() - started) * 1000),
        }
        print(f"Uploaded {spec['filename']}: {body['document']['chunk_count']} chunks", flush=True)
    return sources


def evaluate_research(
    client: httpx.Client,
    sources: dict,
    dataset: dict,
    *,
    prices_per_million: dict[str, Decimal] | None = None,
) -> dict:
    chunks_by_source = {}
    for key, source in sources.items():
        chunks = {}
        offset = 0
        while True:
            response = client.get(
                f"/v1/sources/{source['source_id']}/chunks",
                params={
                    "document_id": source["document_id"],
                    "limit": 200,
                    "offset": offset,
                },
            )
            response.raise_for_status()
            page = response.json()
            chunks.update({chunk["chunk_id"]: chunk for chunk in page["chunks"]})
            offset += len(page["chunks"])
            if offset >= page["total"]:
                break
        chunks_by_source[key] = chunks

    # Resolve annotations against actual parsed chunks before making any generation calls.
    gold = {}
    for case in dataset["cases"]:
        for key, phrases in case["gold_passages"].items():
            matches = {
                chunk_id
                for chunk_id, chunk in chunks_by_source[key].items()
                if any(phrase.casefold() in chunk["text"].casefold() for phrase in phrases)
            }
            if not matches:
                raise ValueError(f"No parsed passage matches the gold label for {case['id']}:{key}")
            gold[(case["id"], key)] = matches

    cases = []
    qrels = []
    rankings = []
    for case in dataset["cases"]:
        selected = {sources[key]["source_id"]: key for key in case["sources"]}
        started = perf_counter()
        response = client.post(
            "/v1/research/runs",
            json={
                "question": case["question"],
                "source_ids": list(selected),
                "mode": case["mode"],
                "retrieval": {"top_k_per_source": 4},
            },
        )
        elapsed_ms = round((perf_counter() - started) * 1000)
        body = response.json()
        checks = {
            name: False
            for name in (
                "schema",
                "expected_status",
                "source_scope",
                "source_coverage",
                "evidence_locators",
                "expected_coverage",
                "candidate_contradictions",
                "qualification",
                "reload",
                "bounded_duration",
                "untrusted_instruction_ignored",
            )
        }
        run = None
        if response.status_code == 200:
            ResearchResponse.model_validate(body)
            checks["schema"] = True
            run_response = client.get(f"/v1/research/runs/{body['run_id']}")
            run_response.raise_for_status()
            run = run_response.json()
            checks["expected_status"] = body["status"] == case["expected_status"]
            checks["source_scope"] = all(
                item["source_id"] in selected
                and item["document_id"] == sources[selected[item["source_id"]]]["document_id"]
                for item in [*body["evidence"], *run["retrieval_results"]]
            )
            checks["source_coverage"] = {
                row["source_id"] for row in body["source_coverage"]
            } == set(selected) and len(body["source_coverage"]) == len(selected)
            checks["expected_coverage"] = {
                selected.get(row["source_id"])
                for row in body["source_coverage"]
                if row["status"] == "used"
            } == set(case["expected_used_sources"])
            checks["evidence_locators"] = all(
                item["source_id"] in selected
                and item["chunk_id"] in chunks_by_source[selected[item["source_id"]]]
                and item["locator"]
                == chunks_by_source[selected[item["source_id"]]][item["chunk_id"]]["locator"]
                and item["excerpt"]
                == chunks_by_source[selected[item["source_id"]]][item["chunk_id"]]["text"]
                for item in body["evidence"]
            )
            checks["candidate_contradictions"] = (
                bool(body["contradictions"]) == case["expect_contradiction"]
            )
            checks["qualification"] = not case.get("expect_qualification") or any(
                link["relationship"] == "qualifies"
                for claim in body["claims"]
                for link in claim["relationships"]
            )
            generated_text = json.dumps(
                {key: body[key] for key in ("summary", "claims", "limitation")}
            )
            checks["untrusted_instruction_ignored"] = not any(
                phrase in generated_text for phrase in case.get("forbidden_answer_text", [])
            )
            reread = client.get(f"/v1/research/runs/{body['run_id']}/result")
            checks["reload"] = reread.status_code == 200 and reread.json() == body
            checks["bounded_duration"] = (
                body["usage"]["duration_ms"] <= run["output"]["timeout_seconds"] * 1000
            )
        elif body.get("error", {}).get("run_id"):
            run_response = client.get(f"/v1/research/runs/{body['error']['run_id']}")
            if run_response.status_code == 200:
                run = run_response.json()
        for key in case["gold_passages"]:
            query_id = f"{case['id']}:{key}"
            qrels.extend(Qrel(query_id, chunk_id, 1) for chunk_id in gold[(case["id"], key)])
            if run:
                rankings.extend(
                    ScoredDoc(query_id, item["chunk_id"], -item["source_rank"])
                    for item in run["retrieval_results"]
                    if item["source_id"] == sources[key]["source_id"]
                )
        cases.append(
            {
                "id": case["id"],
                "question": case["question"],
                "http_status": response.status_code,
                "http_duration_ms": elapsed_ms,
                "checks": checks,
                "passed": all(checks.values()),
                "review_reference": case["review"],
                "response": body,
                "run": run,
            }
        )
        print(
            f"{case['id']}: {'PASS' if all(checks.values()) else 'FAIL'} ({elapsed_ms} ms)",
            flush=True,
        )
    metrics = calc([R @ 1, R @ 4, RR @ 4, nDCG @ 4], qrels, rankings) if qrels else None
    totals = {
        name: sum(
            (case["run"] or case["response"]).get("usage", {}).get(name) or 0 for case in cases
        )
        for name in ("embedding_tokens", "input_tokens", "output_tokens")
    }
    provider_cost = {
        "usd": None,
        "reason": "No verified prices supplied, or a failed request has incomplete usage.",
        "scope": "Query embeddings and generation only; excludes document-ingestion embeddings.",
    }
    if prices_per_million is not None and all(case["http_status"] == 200 for case in cases):
        provider_cost = {
            "usd": str(
                sum(Decimal(count) * prices_per_million[name] for name, count in totals.items())
                / Decimal(1_000_000)
            ),
            "prices_per_million": {name: str(price) for name, price in prices_per_million.items()},
            "scope": provider_cost["scope"],
            "reason": (
                "Estimate using caller-supplied deployment prices, not an Azure billing record."
            ),
        }
    return {
        "dataset_version": dataset["version"],
        "sources": sources,
        "case_count": len(cases),
        "pass_count": sum(case["passed"] for case in cases),
        "check_pass_rates": {
            name: sum(case["checks"][name] for case in cases) / len(cases)
            for name in cases[0]["checks"]
        },
        "retrieval_metrics_per_source": {
            str(key): float(value) for key, value in metrics.aggregated.items()
        }
        if metrics
        else {},
        "query_usage": totals,
        "support_review": (
            "Inspect claims against the stored excerpts and review_reference. Automated checks "
            "do not prove semantic support or general retrieval quality."
        ),
        "provider_cost": provider_cost,
        "cases": cases,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base-url", default="http://127.0.0.1:8000")
    parser.add_argument("--cases", type=Path, default=DEFAULT_CASES)
    parser.add_argument("--output", type=Path, default=Path(".data/mvp2-evaluation.json"))
    parser.add_argument(
        "--reuse-sources", type=Path, help="Reuse source IDs from an earlier report"
    )
    parser.add_argument("--embedding-usd-per-million", type=Decimal)
    parser.add_argument("--input-usd-per-million", type=Decimal)
    parser.add_argument("--output-usd-per-million", type=Decimal)
    args = parser.parse_args()
    prices = {
        "embedding_tokens": args.embedding_usd_per_million,
        "input_tokens": args.input_usd_per_million,
        "output_tokens": args.output_usd_per_million,
    }
    if any(price is not None for price in prices.values()):
        if any(price is None or not price.is_finite() or price < 0 for price in prices.values()):
            parser.error(
                "Provide all three finite, nonnegative deployment prices, or omit all three."
            )
    else:
        prices = None
    dataset = json.loads(args.cases.read_text())
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with httpx.Client(base_url=args.base_url, timeout=240) as client:
        sources = (
            json.loads(args.reuse_sources.read_text())["sources"]
            if args.reuse_sources
            else upload_documents(client, dataset)
        )
        args.output.write_text(json.dumps({"sources": sources}, indent=2))
        report = evaluate_research(client, sources, dataset, prices_per_million=prices)
    args.output.write_text(json.dumps(report, indent=2))
    print(
        json.dumps(
            {key: value for key, value in report.items() if key not in {"cases", "sources"}},
            indent=2,
        )
    )
    if report["pass_count"] != report["case_count"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
